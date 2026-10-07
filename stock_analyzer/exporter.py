"""Excel export of the full step-by-step analysis."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

FILL = {"pass": "C6EFCE", "fail": "F4CCCC", "na": "EDEDED",
        "BUY": "C6EFCE", "WATCH": "FFEB9C", "REJECT": "F4CCCC", "HIGH": "C6EFCE", "MEDIUM": "FFEB9C", "LOW": "F4CCCC"}


def _fill(cell, key):
    if key in FILL:
        cell.fill = PatternFill("solid", fgColor=FILL[key])


def to_excel(a):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    r = a["result"]
    rows = [("Ticker", a["ticker"]), ("Company", a["company"]["name"]), ("Sector", a["company"]["sector"]),
            ("Verdict", r["verdict"]), ("Confidence", r["confidence"]), ("Score", f"{r['score']} / {r['max']}"),
            ("Buy at", r["buy_at"]), ("Watch at", r["watch_at"]), ("", "")]
    rows += [(b["title"], f"{b['points']} / {b['possible']}" if b["possible"] else b["points"]) for b in r["blocks"]]
    rows += [("", "")] + [("Reason", x) for x in r["reasons"]]
    for row in rows:
        ws.append(row)
    for c in ws["B"]:
        _fill(c, str(c.value))
    for c in ws["A"]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 90

    ws = wb.create_sheet("Workings")
    head = ["Step", "Metric", "Result", "Formula", "Working", "Rule", "Outcome", "Points", "Note"]
    ws.append(head)
    for n, s in enumerate(a["steps"], 1):
        for it in s["items"]:
            ws.append([f"{n}. {s['title']}", it["label"], it["display"], it.get("formula"), it.get("working"),
                       it.get("rule"), it["outcome"], it.get("points"), it.get("note")])
            _fill(ws.cell(ws.max_row, 7), it["outcome"])
        for f in (s.get("extra") or {}).get("flags", []) if s["id"] == "flags" else []:
            ws.append([f"{n}. {s['title']}", f["text"], "", "", f["why"], "", "flag", "", ""])
            _fill(ws.cell(ws.max_row, 7), "fail")
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHI", [32, 34, 16, 42, 70, 28, 10, 8, 40]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
